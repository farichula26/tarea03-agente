"""Exporta el Historial a Excel y la tabla percepción-acción a PNG / portapapeles."""
import shutil
import subprocess
import tempfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image, ImageDraw, ImageFont

from agent.ambiente import BORDE as PARED, OSCURO
from agent.reglas import tabla_percepcion_accion
from ui import tema
from ui.simbolos import _rayas

AMARILLO = PatternFill("solid", fgColor="F2DFA4")  # filas ejecutadas (igual que tema.RESALTE)
GRIS = PatternFill("solid", fgColor="D9D9D9")
BORDE = Border(*(Side(style="thin", color="808080"),) * 4)


def exportar_historial(arbol, ruta):
    """Escribe en `ruta` los encabezados y filas del Treeview `arbol` (fila pendiente incluida)."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Historial"
    columnas = arbol["columns"]
    hoja.append([arbol.heading(c, "text") for c in columnas])
    for celda in hoja[1]:
        celda.font, celda.fill, celda.border = Font(bold=True), GRIS, BORDE
        celda.alignment = Alignment(horizontal="center")

    for iid in arbol.get_children():
        valores = list(arbol.item(iid, "values"))
        valores[0] = int(valores[0])  # N como número
        hoja.append(valores)
        hecho = "hecho" in arbol.item(iid, "tags")
        for celda in hoja[hoja.max_row]:
            celda.border = BORDE
            celda.alignment = Alignment(horizontal="center")
            if hecho:
                celda.fill = AMARILLO

    for i, c in enumerate(columnas):
        letra = hoja.cell(row=1, column=i + 1).column_letter
        hoja.column_dimensions[letra].width = 18 if c in ("acciones", "obs") else 8
    hoja.freeze_panes = "A2"
    libro.save(ruta)


# --- Tabla percepción-acción como imagen (para diapositivas) ---

def _fuente(tam, negrita=False):
    """Fuente del sistema (fc-match en Linux, Arial en Windows); si no hay, la que trae Pillow."""
    nombres = ["arialbd.ttf" if negrita else "arial.ttf"]
    if shutil.which("fc-match"):
        ruta = subprocess.run(["fc-match", "-f", "%{file}", "sans:bold" if negrita else "sans"],
                              capture_output=True, text=True).stdout
        nombres.insert(0, ruta)
    for nombre in nombres:
        try:
            return ImageFont.truetype(nombre, tam)
        except OSError:
            pass
    return ImageFont.load_default(size=tam)


def imagen_tabla(escala=2):
    """Los 54 casos en una cuadrícula de 9×6 tarjetas: dibujo del caso, «C# · R#» y la acción."""
    s = escala
    lado, ancho_t, alto_t, margen, cabeza = 34 * s, 150 * s, 150 * s, 20 * s, 56 * s
    img = Image.new("RGB", (2 * margen + 9 * ancho_t, cabeza + 6 * alto_t + 44 * s + margen), "white")
    d = ImageDraw.Draw(img)
    d.text((margen, margen), "Tabla percepción – acción", fill=tema.TEXTO, font=_fuente(22 * s, True))
    grande, chica = _fuente(15 * s, True), _fuente(12 * s)

    def celda(x, y, valor):
        d.rectangle((x, y, x + lado, y + lado), fill=tema.TINTA if valor == OSCURO else tema.TAPETE,
                    outline=tema.REJILLA, width=s)
        if valor == PARED:
            for k in _rayas(lado):
                d.line((x + s, y + k, x + lado - s, y + k), fill=tema.TINTA, width=max(1, s))

    for i, (caso, piso, izq, centro, der, rid, acc) in enumerate(tabla_percepcion_accion()):
        x0, y0 = margen + (i % 9) * ancho_t, cabeza + (i // 9) * alto_t
        d.rectangle((x0, y0, x0 + ancho_t, y0 + alto_t), outline="#B8BDB9", width=s)
        cx, cy = x0 + (ancho_t - 3 * lado) // 2, y0 + 10 * s
        for j, v in enumerate((izq, centro, der)):
            celda(cx + j * lado, cy, v)
        celda(cx + lado, cy + lado, piso)
        d.text((cx + 1.5 * lado, cy + 1.5 * lado), "X", anchor="mm", font=grande,
               fill=tema.TAPETE if piso == OSCURO else tema.TINTA)
        centro_x = x0 + ancho_t / 2
        d.text((centro_x, cy + 2 * lado + 18 * s), f"C{caso} · {rid}", anchor="mm", font=grande, fill=tema.TEXTO)
        d.text((centro_x, cy + 2 * lado + 38 * s), ", ".join(acc).replace("Rotar", "Rot"), anchor="mm",
               font=chica, fill=tema.TEXTO)

    d.text((margen, cabeza + 6 * alto_t + 14 * s),
           "Negro = línea (O), blanco = no-línea (-O), rayado = pared (P); X = agente mirando hacia arriba. "
           "Rot +90 = izquierda, Rot -90 = derecha. Con Contacto, cualquier caso usa R1 (Rot -90).",
           fill=tema.TEXTO_SUAVE, font=chica)
    return img


def guardar_tabla_png(ruta):
    imagen_tabla().save(ruta)


def copiar_tabla_portapapeles():
    """Copia el PNG al portapapeles con wl-copy (Wayland) o xclip (X11). False si no hay ninguno."""
    for cmd in (["wl-copy", "--type", "image/png"], ["xclip", "-selection", "clipboard", "-t", "image/png"]):
        if shutil.which(cmd[0]):
            # Se pasa un archivo (no un pipe) porque wl-copy/xclip quedan en segundo plano sirviendo la imagen
            with tempfile.TemporaryFile() as png:
                imagen_tabla().save(png, "PNG")
                png.seek(0)
                subprocess.run(cmd, stdin=png, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                               check=True, timeout=10)
            return True
    return False
